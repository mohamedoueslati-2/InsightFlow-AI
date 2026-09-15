from io import BytesIO
from pathlib import Path

from docx import Document

try:
    from services.report_presentation.core.config import settings
    from services.report_presentation.services.job_service import create_job, job_directory
except ModuleNotFoundError:
    from backend.services.report_presentation.core.config import settings
    from backend.services.report_presentation.services.job_service import create_job, job_directory


def test_word_report_service_creates_isolated_job(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_dir", tmp_path / "report_jobs")
    document = Document()
    document.add_heading("Quarterly Sales Report", 0)
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph("Revenue increased while margins remained stable.")
    buffer = BytesIO()
    document.save(buffer)

    result = create_job(buffer.getvalue(), "quarterly.docx")
    job_dir = job_directory(result.jobId)

    assert job_dir is not None
    assert result.report.title == "Quarterly Sales Report"
    assert (job_dir / "source" / "report.docx").is_file()
    assert (job_dir / "extraction.json").is_file()
    assert (job_dir / "presentation_input.json").is_file()
    assert (job_dir / "report.md").is_file()
